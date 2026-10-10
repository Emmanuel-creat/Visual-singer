"""Exports : Markdown complet, Word (.docx) et PDF (via LibreOffice si installé)."""
from __future__ import annotations

import re
import shutil
import subprocess
import tempfile
from pathlib import Path

from .llm import undash
from .steps.analyze import fmt
from .steps.documents import soffice
from .store import Project


CRITIQUE = {"urgence": "Urgence", "contre_indication": "Contre-indication", "piege": "Piège", "exception": "Exception",
            "critere": "Critère", "seuil": "Seuil"}


def fiche_markdown(f: dict) -> list[str]:
    """Une fiche : en-tête, 1. définitions, 2. mécanismes, 3. points critiques, 4. auto-évaluation."""
    out = [f"## {f['titre']}", "", f"*Objectif : {f.get('objectif', '')}*", ""]
    out += [f"- **{x}**" for x in f.get("idees_essentielles") or []] + [""]
    if f.get("definitions") or f.get("chiffres") or f.get("classifications"):
        out += ["### 1. Définitions et bases", ""]
        out += [f"- **{d['terme']}** : {d['definition']}" for d in f.get("definitions") or []]
        out += [f"- **{c['valeur']}** : {c['signification']}" for c in f.get("chiffres") or []]
        for c in f.get("classifications") or []:
            out += [f"- **{c['titre']}** : " + " ; ".join(c.get("elements") or [])]
        out.append("")
    if f.get("mecanismes"):
        out += ["### 2. Mécanismes et raisonnement", ""]
        for m in f["mecanismes"]:
            out += [f"**{m['titre']}** : " + " → ".join(m.get("etapes") or []), ""]
    if f.get("points_critiques") or f.get("a_ne_pas_oublier"):
        out += ["### 3. Points critiques et exceptions", ""]
        out += [f"- **À ne pas oublier** : {x}" for x in f.get("a_ne_pas_oublier") or []]
        out += [f"- **{CRITIQUE.get(c['type'], c['type'])}** : {c['texte']}" for c in f.get("points_critiques") or []]
        out.append("")
    comp = f.get("comparatif")
    if comp:
        out += [f"**{comp['titre']}**", "", "| Critère | " + " | ".join(comp["colonnes"]) + " |",
                "|---|" + "---|" * len(comp["colonnes"])]
        out += ["| " + " | ".join([l["critere"], *[(l["valeurs"] + [""] * len(comp["colonnes"]))[i] for i in range(len(comp["colonnes"]))]]) + " |"
                for l in comp["lignes"]]
        out.append("")
    if f.get("auto_evaluation"):
        out += ["### 4. Auto-évaluation", ""]
        out += [f"{i}. {q['question']}" for i, q in enumerate(f["auto_evaluation"], 1)] + [""]
        out += ["*Réponses* : " + " ; ".join(f"{i}. {q['reponse']}" for i, q in enumerate(f["auto_evaluation"], 1)), ""]
    return out


def build_markdown(project: Project, with_transcript: bool = True) -> str:
    meta = project.meta()
    a = project.result("analyze") or {}
    course = a.get("course", {})
    paras = (project.result("paragraphs") or {}).get("paragraphs", [])
    speakers = (project.result("paragraphs") or {}).get("speakers", [])
    out = [f"# {course.get('titre_cours') or meta.get('title', 'Cours')}", ""]
    if course.get("resume_court"):
        out += ["## Résumé court", "", course["resume_court"], ""]
    if a.get("synthesis"):
        out += ["# Synthèse finale (audio + support)", "", a["synthesis"].strip(), ""]
    fiches = (project.result("fiches") or {}).get("fiches", [])
    if fiches:
        out += ["# Fiches de révision", ""]
        for f in fiches:
            out += fiche_markdown(f)
    if course.get("plan"):
        out += ["## Plan du cours", ""]
        secs = a.get("sections", [])
        for i, p in enumerate(course["plan"], 1):
            t0 = min((secs[s]["start"] for s in p.get("sections", []) if s < len(secs)), default=None)
            out.append(f"{i}. **{p['titre']}**" + (f" ({fmt(t0)})" if t0 is not None else ""))
            for sp in p.get("sous_parties", []):
                out.append(f"   - {sp}")
        out.append("")
    if course.get("resume_detaille"):
        out += ["## Résumé détaillé", "", course["resume_detaille"].strip(), ""]
    if a.get("glossary"):
        out += ["## Notions et définitions", ""]
        out += [f"- **{g['terme']}** ({fmt(g.get('t') or 0)}) : {g['definition']}" for g in a["glossary"]]
        out.append("")
    if a.get("formulas"):
        out += ["## Formules et règles", ""]
        out += [f"- `{f['formule']}` : {f['explication']} ({fmt(f.get('t') or 0)})" for f in a["formulas"]]
        out.append("")
    if a.get("examples"):
        out += ["## Exemples", ""] + [f"- {e['exemple']} ({fmt(e.get('t') or 0)})" for e in a["examples"]] + [""]
    if a.get("keep"):
        out += ["## À retenir", ""] + [f"- {k['point']}" for k in a["keep"]] + [""]
    cross = a.get("cross") or {}
    if cross.get("ecarts") or cross.get("support_non_explique") or cross.get("oral_absent_support"):
        out += ["## Croisement oral / support", ""]
        labels = {"contradiction": "Contradiction", "difference": "Différence", "ajout_oral": "Ajout à l'oral", "non_explique": "Non expliqué"}
        for e in cross.get("ecarts", []):
            out.append(f"- **{labels.get(e['type'], e['type'])}** ({fmt(e.get('t') or 0)}"
                       + (f", diapositive {e['diapositive']}" if e.get("diapositive") else "") + f") : {e['detail']}")
        if cross.get("oral_absent_support"):
            out += ["", "### Notions expliquées à l'oral, absentes du support", ""]
            out += [f"- {g['terme']} ({fmt(g.get('t') or 0)})" for g in cross["oral_absent_support"]]
        if cross.get("support_non_explique"):
            out += ["", "### Éléments du support peu ou pas expliqués à l'oral", ""]
            out += [f"- {c['line']} ({c.get('doc_name', '')}, {c.get('kind', 'page')} {c.get('n')})" for c in cross["support_non_explique"][:60]]
        out.append("")
    if a.get("qa"):
        out += ["## Questions de révision", ""]
        for i, q in enumerate(a["qa"], 1):
            out += [f"**{i}. {q['question']}**", "", f"> {q['reponse']}", ""]
    if a.get("sections"):
        out += ["## Détail par section", ""]
        for r in a["sections"]:
            out += [f"### {r['section'] + 1}. {r['titre']} ({fmt(r['start'])} – {fmt(r['end'])})", "", r["resume"], ""]
            out += [f"- {p}" for p in r.get("points_cles", [])] + [""]
    if with_transcript and paras:
        out += ["## Transcription", ""]
        for p in paras:
            spk = f"**{speakers[p['speaker']]}** " if len(speakers) > 1 else ""
            out += [f"[{fmt(p['start'])}] {spk}{p['text']}", ""]
    return "\n".join(out)


def _add_runs(par, text: str) -> None:
    """Gras **…** et code `…` minimalistes."""
    for tok in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            par.add_run(tok[2:-2]).bold = True
        elif tok.startswith("`") and tok.endswith("`"):
            r = par.add_run(tok[1:-1])
            r.font.name = "Consolas"
        else:
            par.add_run(tok)


def markdown_to_docx(md: str, path: Path) -> None:
    import docx
    from docx.shared import Pt
    d = docx.Document()
    d.styles["Normal"].font.name = "Calibri"
    d.styles["Normal"].font.size = Pt(11)
    lines = md.splitlines()
    i = 0
    while i < len(lines):
        s = lines[i].rstrip()
        i += 1
        if not s:
            continue
        if s.startswith("|"):
            # Tableau Markdown (tableau comparatif des fiches) → vrai tableau Word.
            rows = [s]
            while i < len(lines) and lines[i].startswith("|"):
                rows.append(lines[i].rstrip())
                i += 1
            cells = [[c.strip() for c in r.strip("|").split("|")] for r in rows if not re.fullmatch(r"\|[\s:|-]+\|?", r)]
            ncol = max(len(r) for r in cells)
            t = d.add_table(rows=len(cells), cols=ncol)
            t.style = "Light Grid Accent 1"
            for r, row in enumerate(cells):
                for c in range(ncol):
                    _add_runs(t.cell(r, c).paragraphs[0], row[c] if c < len(row) else "")
            continue
        m = re.match(r"^(#{1,4})\s+(.*)", s)
        if m:
            d.add_heading(m.group(2), level=min(4, len(m.group(1))))
        elif re.match(r"^\s*[-*]\s+", s):
            indent = (len(s) - len(s.lstrip())) // 2
            p = d.add_paragraph(style="List Bullet" if indent == 0 else "List Bullet 2")
            _add_runs(p, re.sub(r"^\s*[-*]\s+", "", s))
        elif re.match(r"^\s*\d+\.\s+", s):
            p = d.add_paragraph(style="List Number")
            _add_runs(p, re.sub(r"^\s*\d+\.\s+", "", s))
        elif s.startswith(">"):
            p = d.add_paragraph(style="Intense Quote")
            _add_runs(p, s.lstrip("> "))
        else:
            _add_runs(d.add_paragraph(), s)
    d.save(str(path))


def export(project: Project, fmt_: str, with_transcript: bool = True) -> Path:
    md = undash(build_markdown(project, with_transcript))
    title = re.sub(r"[^\w\- ]+", "", project.meta().get("title", "cours")).strip().replace(" ", "_")[:60] or "cours"
    project.exports.mkdir(exist_ok=True)
    if fmt_ == "md":
        p = project.exports / f"{title}.md"
        p.write_text(md, encoding="utf-8")
        return p
    docx_path = project.exports / f"{title}.docx"
    markdown_to_docx(md, docx_path)
    if fmt_ == "docx":
        return docx_path
    if fmt_ == "pdf":
        exe = soffice()
        if not exe:
            raise RuntimeError("L'export PDF nécessite LibreOffice (export Word disponible).")
        tmp = Path(tempfile.mkdtemp(prefix="coursia-pdf-"))
        subprocess.run([exe, "--headless", "--norestore", "--convert-to", "pdf", "--outdir", str(tmp), str(docx_path)],
                       capture_output=True, timeout=600)
        pdfs = list(tmp.glob("*.pdf"))
        if not pdfs:
            raise RuntimeError("La conversion PDF a échoué")
        out = project.exports / f"{title}.pdf"
        shutil.move(str(pdfs[0]), out)
        shutil.rmtree(tmp, ignore_errors=True)
        return out
    raise ValueError(fmt_)
