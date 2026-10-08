"""Extraction des supports de cours : PDF, PowerPoint, Word, texte.

Chaque document devient une liste d'« unités » (diapositive, page ou section) :
  { id, doc, kind, n, title, text, notes, tables: [markdown], ocr, image }

- PDF : texte par blocs, titres (plus grande police), tableaux, miniature.
- PPTX : titre, puces avec niveaux, tableaux, notes de l'orateur ; miniatures via
  LibreOffice (conversion PDF headless) si installé.
- DOCX : sections découpées sur les titres, tableaux.
- OCR (EasyOCR, GPU) quand une page/diapositive ne contient presque pas de texte
  (scan, diapositive-image).

Le vocabulaire du support (titres, sigles, termes techniques récurrents) est extrait
pour guider la transcription.
"""
from __future__ import annotations

import re
import shutil
import subprocess
import tempfile
from collections import Counter
from pathlib import Path

from ..store import Project, cache_key, write_json

VERSION = "docs-v1"
OCR_MIN_CHARS = 40
THUMB_WIDTH = 1280

_ocr_reader = None


def _ocr(img_bytes: bytes) -> str:
    global _ocr_reader
    if _ocr_reader is None:
        import easyocr
        from ..hardware import detect
        _ocr_reader = easyocr.Reader(["fr", "en"], gpu=detect().torch_cuda, verbose=False)
    lines = _ocr_reader.readtext(img_bytes, detail=0, paragraph=True)
    return "\n".join(lines)


def release() -> None:
    global _ocr_reader
    _ocr_reader = None


def soffice() -> str | None:
    for cand in (shutil.which("soffice"), r"C:\Program Files\LibreOffice\program\soffice.exe",
                 r"C:\Program Files (x86)\LibreOffice\program\soffice.exe", "/usr/bin/soffice",
                 "/Applications/LibreOffice.app/Contents/MacOS/soffice"):
        if cand and Path(cand).exists():
            return cand
    return None


def to_pdf(src: Path) -> Path | None:
    """Conversion PPTX/DOCX → PDF avec LibreOffice (pour les miniatures et l'OCR)."""
    exe = soffice()
    if not exe:
        return None
    from ..config import MODELS_DIR
    out = Path(tempfile.mkdtemp(prefix="coursia-"))
    # Profil LibreOffice dédié (persistant) : indépendant d'un LibreOffice déjà ouvert.
    # Au tout premier lancement, LibreOffice initialise le profil sans convertir : on réessaie.
    profile = (MODELS_DIR / "libreoffice-profile").resolve().as_uri()
    for attempt in range(3):
        try:
            r = subprocess.run([exe, f"-env:UserInstallation={profile}", "--headless", "--norestore", "--nolockcheck",
                                "--convert-to", "pdf", "--outdir", str(out), str(src)],
                               capture_output=True, timeout=600, text=True, errors="replace")
        except Exception as e:
            print(f"[LibreOffice] {e}", flush=True)
            return None
        pdfs = list(out.glob("*.pdf"))
        if pdfs:
            return pdfs[0]
    print(f"[LibreOffice] conversion impossible : {r.stdout[-300:]} {r.stderr[-300:]}", flush=True)
    return None


def _table_md(rows: list[list[str]]) -> str:
    rows = [[(c or "").replace("\n", " ").strip() for c in r] for r in rows if any((c or "").strip() for c in r)]
    if not rows:
        return ""
    w = max(len(r) for r in rows)
    rows = [r + [""] * (w - len(r)) for r in rows]
    head = "| " + " | ".join(rows[0]) + " |"
    sep = "| " + " | ".join(["---"] * w) + " |"
    body = ["| " + " | ".join(r) + " |" for r in rows[1:]]
    return "\n".join([head, sep, *body])


def _render(page, path: Path, width: int = THUMB_WIDTH) -> bytes:
    import pymupdf as fitz
    zoom = width / max(1, page.rect.width)
    pix = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom), alpha=False)
    path.parent.mkdir(parents=True, exist_ok=True)
    pix.save(str(path))
    return pix.tobytes("png")


# ------------------------------------------------------------------ PDF
def extract_pdf(path: Path, doc_id: str, outdir: Path, ctx, kind: str = "page", base_units: list | None = None) -> list[dict]:
    import pymupdf as fitz
    units = []
    with fitz.open(str(path)) as pdf:
        n = len(pdf)
        for i, page in enumerate(pdf):
            ctx.progress(i / max(1, n))
            d = page.get_text("dict")
            spans = [(s["size"], s["text"].strip()) for b in d["blocks"] if b.get("type") == 0
                     for l in b["lines"] for s in l["spans"] if s["text"].strip()]
            text = page.get_text("text").strip()
            title = ""
            if spans:
                big = max(s[0] for s in spans)
                title = " ".join(t for sz, t in spans if sz >= big - 0.5)[:160]
            tables = []
            try:
                for t in page.find_tables().tables:
                    md = _table_md(t.extract())
                    if md:
                        tables.append(md)
            except Exception:
                pass
            img_rel = f"docs/{doc_id}/p{i + 1:03d}.png"
            png = _render(page, outdir / f"p{i + 1:03d}.png")
            ocr = False
            if len(text) < OCR_MIN_CHARS:
                ocr_text = _ocr(png)
                if len(ocr_text) > len(text):
                    text, ocr = ocr_text, True
                    title = title or ocr_text.split("\n")[0][:160]
            if base_units and i < len(base_units):
                # Diapositive PowerPoint : on garde le texte structuré, on ajoute l'image (et l'OCR si vide).
                u = base_units[i]
                u["image"] = img_rel
                if ocr and len(u["text"]) < OCR_MIN_CHARS:
                    u["text"], u["ocr"] = text, True
                continue
            units.append({"id": f"{doc_id}-{i + 1}", "doc": doc_id, "kind": kind, "n": i + 1, "title": title,
                          "text": text, "notes": "", "tables": tables, "ocr": ocr, "image": img_rel})
    return base_units if base_units else units


# ------------------------------------------------------------------ PPTX
def extract_pptx(path: Path, doc_id: str, outdir: Path, ctx) -> list[dict]:
    from pptx import Presentation
    from pptx.util import Emu  # noqa: F401
    prs = Presentation(str(path))
    units = []
    for i, slide in enumerate(prs.slides):
        title, lines, tables = "", [], []
        for shape in slide.shapes:
            if shape.has_text_frame:
                is_title = shape.is_placeholder and shape.placeholder_format.type is not None and "TITLE" in str(shape.placeholder_format.type)
                for p in shape.text_frame.paragraphs:
                    t = "".join(r.text for r in p.runs).strip()
                    if not t:
                        continue
                    if is_title and not title:
                        title = t
                    else:
                        lines.append(("  " * p.level) + "• " + t)
            if getattr(shape, "has_table", False) and shape.has_table:
                rows = [[c.text for c in r.cells] for r in shape.table.rows]
                md = _table_md(rows)
                if md:
                    tables.append(md)
        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        units.append({"id": f"{doc_id}-{i + 1}", "doc": doc_id, "kind": "slide", "n": i + 1, "title": title or (lines[0][2:80] if lines else f"Diapositive {i + 1}"),
                      "text": "\n".join(lines), "notes": notes, "tables": tables, "ocr": False, "image": None})
    ctx.log(f"{len(units)} diapositives lues ; conversion en images via LibreOffice…")
    pdf = to_pdf(path)
    if pdf:
        units = extract_pdf(pdf, doc_id, outdir, ctx, kind="slide", base_units=units)
        shutil.rmtree(pdf.parent, ignore_errors=True)
    else:
        ctx.log("Conversion LibreOffice impossible : pas de miniatures pour les diapositives")
    return units


# ------------------------------------------------------------------ DOCX
def extract_docx(path: Path, doc_id: str, ctx) -> list[dict]:
    import docx
    d = docx.Document(str(path))
    units, cur = [], {"title": "Début du document", "lines": []}

    def flush():
        if cur["lines"]:
            n = len(units) + 1
            units.append({"id": f"{doc_id}-{n}", "doc": doc_id, "kind": "section", "n": n, "title": cur["title"],
                          "text": "\n".join(cur["lines"]), "notes": "", "tables": [], "ocr": False, "image": None})

    for p in d.paragraphs:
        t = p.text.strip()
        if not t:
            continue
        style = (p.style.name or "").lower()
        if style.startswith("heading") or style.startswith("titre") or style == "title":
            flush()
            cur = {"title": t, "lines": []}
        else:
            cur["lines"].append(t)
    flush()
    for tb in d.tables:
        md = _table_md([[c.text for c in r.cells] for r in tb.rows])
        if md and units:
            units[-1]["tables"].append(md)
    return units


def extract_text(path: Path, doc_id: str) -> list[dict]:
    raw = path.read_text(encoding="utf-8", errors="replace")
    parts = re.split(r"\n(?=#{1,3} )", raw)
    units = []
    for i, part in enumerate(p for p in parts if p.strip()):
        lines = part.strip().splitlines()
        units.append({"id": f"{doc_id}-{i + 1}", "doc": doc_id, "kind": "section", "n": i + 1,
                      "title": lines[0].lstrip("# ")[:160], "text": "\n".join(lines[1:]), "notes": "", "tables": [], "ocr": False, "image": None})
    return units


# ------------------------------------------------------------------ vocabulaire
STOP = set("""le la les un une des du de d l et ou à au aux en dans sur pour par avec sans ce cet cette ces se sa son ses leur leurs
nous vous ils elles il elle on ne pas plus moins est sont été être avoir a ont qui que quoi dont où comme mais donc car si tout tous
the of and to in for on with by is are be this that from an as at or not it its""".split())


def vocabulary(units: list[dict], limit: int = 60) -> list[str]:
    """Termes du support susceptibles d'être mal transcrits : sigles, noms propres, termes techniques."""
    counts = Counter()
    for u in units:
        blob = " ".join([u.get("title", ""), u.get("text", ""), u.get("notes", "")])
        for w in re.findall(r"[A-Za-zÀ-ÖØ-öø-ÿ][\w\-’']{2,}", blob):
            lw = w.lower()
            if lw in STOP or w.isdigit():
                continue
            score = 1
            if w.isupper() and len(w) >= 2:
                score = 3                       # sigle
            elif w[0].isupper():
                score = 2                       # nom propre / notion
            elif len(w) >= 9:
                score = 1.5                     # terme technique long
            counts[w] += score
        for t in [u.get("title", "")]:
            if 3 <= len(t) <= 60:
                counts[t.strip()] += 2
    terms = [t for t, c in counts.most_common(limit * 2) if c >= 2]
    seen, out = set(), []
    for t in terms:
        if t.lower() not in seen:
            seen.add(t.lower())
            out.append(t)
        if len(out) >= limit:
            break
    return out


# ------------------------------------------------------------------ étape
def key(project: Project) -> str:
    docs = project.meta().get("docs") or []
    return cache_key(VERSION, [(d["id"], d["fingerprint"]) for d in docs])


def run(project: Project, ctx) -> dict:
    docs = project.meta().get("docs") or []
    k = key(project)
    cached = project.cached("documents", k)
    if cached is not None:
        return cached
    all_units = []
    for j, d in enumerate(docs):
        src = project.originals / d["file"]
        outdir = project.docs_dir / d["id"]
        outdir.mkdir(parents=True, exist_ok=True)
        ext = src.suffix.lower()
        ctx.log(f"Lecture de {d['name']}")
        sub = ctx.sub(j / max(1, len(docs)), (j + 1) / max(1, len(docs)))
        if ext == ".pdf":
            units = extract_pdf(src, d["id"], outdir, sub)
        elif ext in (".pptx", ".ppsx", ".pptm"):
            units = extract_pptx(src, d["id"], outdir, sub)
        elif ext in (".ppt", ".odp", ".key"):
            pdf = to_pdf(src)
            if not pdf:
                raise RuntimeError(f"{d['name']} : format ancien, LibreOffice est nécessaire pour le lire")
            units = extract_pdf(pdf, d["id"], outdir, sub, kind="slide")
        elif ext in (".docx", ".odt", ".doc"):
            if ext != ".docx":
                pdf = to_pdf(src)
                if not pdf:
                    raise RuntimeError(f"{d['name']} : LibreOffice est nécessaire pour ce format")
                units = extract_pdf(pdf, d["id"], outdir, sub)
            else:
                units = extract_docx(src, d["id"], sub)
        elif ext in (".txt", ".md"):
            units = extract_text(src, d["id"])
        else:
            raise RuntimeError(f"Format de document non pris en charge : {d['name']}")
        for u in units:
            u["doc_name"] = d["name"]
        write_json(outdir / "units.json", units)
        n_ocr = sum(1 for u in units if u.get("ocr"))
        ctx.log(f"{d['name']} : {len(units)} {'diapositives' if units and units[0]['kind'] == 'slide' else 'pages/sections'}"
                + (f", dont {n_ocr} lues par OCR" if n_ocr else ""))
        all_units.extend(units)
    release()
    data = {"units": all_units, "vocabulary": vocabulary(all_units)}
    project.put("documents", k, data)
    return data
